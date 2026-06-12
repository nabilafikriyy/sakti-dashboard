from flask import Flask, jsonify, render_template, request
from flask_cors import CORS
import openpyxl, os, sqlite3, datetime

app = Flask(__name__)
CORS(app)
DATA_DIR = os.path.join(os.path.dirname(__file__), 'data')
DB_PATH  = os.path.join(os.path.dirname(__file__), 'sakti.db')

FILE_DBD      = 'analisis_dbd_per_kelompok_umur_updated.xlsx'
FILE_STUNTING = 'gizi_balita_psg_2023_2025_updated.xlsx'
FILE_CAMPAK   = 'template_campak_sleman_2023-2025.xlsx'

# ── DATABASE ADAPTER (DUAL MODE: LOCAL SQLITE & VERCEL POSTGRES) ──
class DBWrapper:
    def __init__(self, conn, is_postgres):
        self.conn = conn
        self.is_postgres = is_postgres
        
    def execute(self, query, params=()):
        if self.is_postgres:
            # Menyesuaikan query SQLite agar otomatis kompatibel dengan PostgreSQL Cloud
            query = query.replace('INTEGER PRIMARY KEY AUTOINCREMENT', 'SERIAL PRIMARY KEY')
            query = query.replace('?', '%s')
            
            is_insert = query.strip().upper().startswith("INSERT")
            if is_insert:
                query = query.strip().rstrip(';') + " RETURNING id"
                
            cur = self.conn.cursor()
            cur.execute(query, params)
            
            last_id = None
            if is_insert:
                res = cur.fetchone()
                if res:
                    last_id = res[0]
                    
            class CursorAdapter:
                def __init__(self, cursor, last_id):
                    self.cursor = cursor
                    self.lastrowid = last_id
                def fetchall(self):
                    return self.cursor.fetchall()
                def fetchone(self):
                    return self.cursor.fetchone()
                    
            return CursorAdapter(cur, last_id)
        else:
            return self.conn.execute(query, params)
            
    def commit(self):
        self.conn.commit()
        
    def close(self):
        self.conn.close()
        
    def __enter__(self):
        return self
        
    def __exit__(self, exc_type, exc_val, exc_tb):
        if exc_type:
            self.conn.rollback()
        else:
            self.conn.commit()
        self.conn.close()

def get_db():
    postgres_url = os.environ.get("DATABASE_URL") or os.environ.get("POSTGRES_URL")
    if postgres_url:
        import psycopg2
        from psycopg2.extras import DictCursor
        if postgres_url.startswith("postgres://"):
            postgres_url = postgres_url.replace("postgres://", "postgresql://", 1)
        conn = psycopg2.connect(postgres_url, cursor_factory=DictCursor)
        return DBWrapper(conn, is_postgres=True)
    else:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        return DBWrapper(conn, is_postgres=False)

def init_db():
    with get_db() as conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS catatan (
                id        INTEGER PRIMARY KEY AUTOINCREMENT,
                judul     TEXT    NOT NULL,
                isi       TEXT    NOT NULL,
                prioritas TEXT    DEFAULT 'Sedang',
                bidang    TEXT    DEFAULT 'Umum',
                penulis   TEXT    DEFAULT 'Kepala Dinas',
                dibuat    TEXT    NOT NULL,
                diubah    TEXT
            )
        ''')
        conn.commit()

init_db()

def load_wb(f): return openpyxl.load_workbook(os.path.join(DATA_DIR, f), data_only=True)
def si(v): return int(v or 0)
def sf(v): return round(float(v or 0), 2)
def avg(lst): return round(sum(lst)/len(lst), 2) if lst else 0

# ═══════════════════════════════════════════════════════════════
# DBD
# ═══════════════════════════════════════════════════════════════
def _read_dbd_year(tahun):
    """Baca satu sheet tahun DBD, kembalikan dict per puskesmas."""
    ws = load_wb(FILE_DBD)[str(tahun)]
    rows = list(ws.iter_rows(values_only=True))
    result = {}
    for row in rows[5:30]:
        if row[2] is None: continue
        pusk = str(row[2]).strip()
        kap  = str(row[1]).strip() if row[1] else '-'
        ages = [si(row[i]) for i in range(3, 11)]
        total = sum(ages)
        result[pusk] = {
            'kapanewon': kap,
            'ages': ages,
            'total': total,
            'laki': si(row[12]),
            'perempuan': si(row[13]),
            'meninggal': si(row[14]),
            'cfr': sf(row[15]),
        }
    return result

@app.route('/api/dbd/rekapitulasi')
def dbd_rekapitulasi():
    data = {}
    for yr in [2023, 2024, 2025]:
        data[yr] = _read_dbd_year(yr)

    ws_order = load_wb(FILE_DBD)['2023']
    rows_order = list(ws_order.iter_rows(values_only=True))
    puskesmas = [str(r[2]).strip() for r in rows_order[5:30] if r[2]]

    y23 = [data[2023].get(p, {}).get('total', 0) for p in puskesmas]
    y24 = [data[2024].get(p, {}).get('total', 0) for p in puskesmas]
    y25 = [data[2025].get(p, {}).get('total', 0) for p in puskesmas]

    return jsonify({
        'puskesmas': puskesmas,
        'tahun_2023': y23, 'tahun_2024': y24, 'tahun_2025': y25,
        'total_2023': sum(y23), 'total_2024': sum(y24), 'total_2025': sum(y25),
        'laki_2023':  sum(data[2023].get(p,{}).get('laki',0) for p in puskesmas),
        'perempuan_2023': sum(data[2023].get(p,{}).get('perempuan',0) for p in puskesmas),
        'laki_2024':  sum(data[2024].get(p,{}).get('laki',0) for p in puskesmas),
        'perempuan_2024': sum(data[2024].get(p,{}).get('perempuan',0) for p in puskesmas),
        'laki_2025':  sum(data[2025].get(p,{}).get('laki',0) for p in puskesmas),
        'perempuan_2025': sum(data[2025].get(p,{}).get('perempuan',0) for p in puskesmas),
    })

@app.route('/api/dbd/kelompok_umur')
def dbd_kelompok_umur():
    age_labels = ['<1','1-4','5-9','10-14','15-19','20-44','45-59','≥60']
    result = {}
    for yr in [2023, 2024, 2025]:
        d = _read_dbd_year(yr)
        totals = [0]*8
        for pusk_data in d.values():
            for i, v in enumerate(pusk_data['ages']):
                totals[i] += v
        result[yr] = totals
    return jsonify({
        'age_groups': age_labels,
        'data_2023': result[2023],
        'data_2024': result[2024],
        'data_2025': result[2025],
    })

@app.route('/api/dbd/gender')
def dbd_gender():
    result = {}
    for yr in [2023, 2024, 2025]:
        d = _read_dbd_year(yr)
        result[yr] = {
            'laki': sum(v['laki'] for v in d.values()),
            'perempuan': sum(v['perempuan'] for v in d.values()),
        }
    return jsonify({
        'laki_2023': result[2023]['laki'], 'perempuan_2023': result[2023]['perempuan'],
        'laki_2024': result[2024]['laki'], 'perempuan_2024': result[2024]['perempuan'],
        'laki_2025': result[2025]['laki'], 'perempuan_2025': result[2025]['perempuan'],
    })

@app.route('/api/dbd/pertahun/<int:tahun>')
def dbd_pertahun(tahun):
    ws = load_wb(FILE_DBD)[str(tahun)]
    rows = list(ws.iter_rows(values_only=True))
    age_labels = ['<1','1-4','5-9','10-14','15-19','20-44','45-59','≥60']
    puskesmas, kapanewon, total, meninggal, cfr = [], [], [], [], []
    laki_list, perempuan_list = [], []
    age_data = {a: [] for a in age_labels}
    kapanewon_agg = {}

    for row in rows[5:30]:
        if row[2] is None: continue
        kap  = str(row[1]).strip() if row[1] else '-'
        pusk = str(row[2]).strip()
        ages = [si(row[i]) for i in range(3, 11)]
        tot  = sum(ages)
        men  = si(row[14])
        c    = sf(row[15])
        l    = si(row[12])
        p    = si(row[13])

        puskesmas.append(pusk); kapanewon.append(kap)
        total.append(tot); meninggal.append(men); cfr.append(c)
        laki_list.append(l); perempuan_list.append(p)
        for i, a in enumerate(age_labels): age_data[a].append(ages[i])
        kapanewon_agg[kap] = kapanewon_agg.get(kap, 0) + tot

    age_total = [sum(age_data[a]) for a in age_labels]

    return jsonify({
        'puskesmas': puskesmas, 'kapanewon': kapanewon,
        'total': total, 'meninggal': meninggal, 'cfr': cfr,
        'laki': laki_list, 'perempuan': perempuan_list,
        'age_data': age_data, 'age_labels': age_labels, 'age_total': age_total,
        'kapanewon_labels': list(kapanewon_agg.keys()),
        'kapanewon_total': list(kapanewon_agg.values()),
        'grand_total': sum(total), 'total_meninggal': sum(meninggal),
    })

# ═══════════════════════════════════════════════════════════════
# STUNTING
# ═══════════════════════════════════════════════════════════════
def _read_stunting_year(tahun):
    ws = load_wb(FILE_STUNTING)[str(tahun)]
    rows = list(ws.iter_rows(values_only=True))
    puskesmas, sasaran, dipantau = [], [], []
    sp_n, sp_pct, p_n, p_pct = [], [], [], []
    normal_n, normal_pct, tinggi_n, tinggi_pct = [], [], [], []
    stunted_n, stunted_pct, stunting_n, stunting_pct, cakupan = [], [], [], [], []

    for row in rows[6:31]:
        if row[1] is None or str(row[1]).strip().upper() in ('JUMLAH','TOTAL',''): continue
        sas = si(row[2]); dip = si(row[3])
        puskesmas.append(str(row[1]).strip())
        sasaran.append(sas); dipantau.append(dip)
        sp_n.append(si(row[4]));    sp_pct.append(sf(row[5]))
        p_n.append(si(row[6]));     p_pct.append(sf(row[7]))
        normal_n.append(si(row[8]));  normal_pct.append(sf(row[9]))
        tinggi_n.append(si(row[10])); tinggi_pct.append(sf(row[11]))
        stunted_n.append(si(row[12]));  stunted_pct.append(sf(row[13]))
        stunting_n.append(si(row[14])); stunting_pct.append(sf(row[15]))
        cak = round(dip / sas * 100, 2) if sas > 0 else 0
        cakupan.append(cak)

    return {
        'puskesmas': puskesmas, 'sasaran': sasaran, 'dipantau': dipantau,
        'sangat_pendek_n': sp_n, 'sangat_pendek_pct': sp_pct,
        'pendek_n': p_n, 'pendek_pct': p_pct,
        'normal_n': normal_n, 'normal_pct': normal_pct,
        'tinggi_n': tinggi_n, 'tinggi_pct': tinggi_pct,
        'stunted_n': stunted_n, 'stunted_pct': stunted_pct,
        'stunting_n': stunting_n, 'stunting_pct': stunting_pct,
        'cakupan': cakupan,
        'kab_sp': sum(sp_n), 'kab_p': sum(p_n),
        'kab_normal': sum(normal_n), 'kab_tinggi': sum(tinggi_n),
        'total_sasaran': sum(sasaran), 'total_dipantau': sum(dipantau),
        'total_stunted': sum(stunted_n), 'total_stunting': sum(stunting_n),
        'avg_cakupan': avg(cakupan),
    }

@app.route('/api/stunting/rekapitulasi')
def stunting_rekapitulasi():
    d23 = _read_stunting_year(2023)
    d24 = _read_stunting_year(2024)
    d25 = _read_stunting_year(2025)

    puskesmas = d23['puskesmas']

    def get_pct(d, key, p):
        if p in d['puskesmas']:
            i = d['puskesmas'].index(p)
            return d[key][i]
        return 0.0

    s23 = [get_pct(d23, 'stunting_pct', p) for p in puskesmas]
    s24 = [get_pct(d24, 'stunting_pct', p) for p in puskesmas]
    s25 = [get_pct(d25, 'stunting_pct', p) for p in puskesmas]

    return jsonify({
        'puskesmas': puskesmas,
        'stunting_2023': s23, 'stunting_2024': s24, 'stunting_2025': s25,
        'avg_2023': avg(s23), 'avg_2024': avg(s24), 'avg_2025': avg(s25),
    })

@app.route('/api/stunting/pertahun/<int:tahun>')
def stunting_pertahun(tahun):
    d = _read_stunting_year(tahun)
    return jsonify(d)

# ═══════════════════════════════════════════════════════════════
# CAMPAK
# ═══════════════════════════════════════════════════════════════
@app.route('/api/campak/rekapitulasi')
def campak_rekapitulasi():
    ws = load_wb(FILE_CAMPAK)['Rekapitulasi']
    rows = list(ws.iter_rows(values_only=True))
    puskesmas=[]
    sp23,po23,r23,sp24,po24,r24,sp25,po25,r25=[],[],[],[],[],[],[],[],[]
    for row in rows[5:30]:
        if row[1] is None: continue
        puskesmas.append(str(row[1]).strip())
        sp23.append(si(row[2])); po23.append(si(row[3])); r23.append(sf(row[4]))
        sp24.append(si(row[5])); po24.append(si(row[6])); r24.append(sf(row[7]))
        sp25.append(si(row[8])); po25.append(si(row[9])); r25.append(sf(row[10] or 0))
    return jsonify({'puskesmas':puskesmas,
        'suspek_2023':sp23,'positif_2023':po23,'rate_2023':r23,
        'suspek_2024':sp24,'positif_2024':po24,'rate_2024':r24,
        'suspek_2025':sp25,'positif_2025':po25,'rate_2025':r25,
        'total_suspek_2023':sum(sp23),'total_positif_2023':sum(po23),
        'total_suspek_2024':sum(sp24),'total_positif_2024':sum(po24),
        'total_suspek_2025':sum(sp25),'total_positif_2025':sum(po25)})

@app.route('/api/campak/pertahun/<int:tahun>')
def campak_pertahun(tahun):
    ws = load_wb(FILE_CAMPAK)[f'Kasus {tahun}']
    rows = list(ws.iter_rows(values_only=True))
    age_labels = ['<1','1-4','5-9','10-14','15-19','≥20']
    puskesmas,kapanewon,total_suspek,total_positif,rate=[],[],[],[],[]
    kap_agg_sp,kap_agg_po={},{}
    age_sp={a:[] for a in age_labels}
    age_po={a:[] for a in age_labels}
    for row in rows[6:31]:
        if row[2] is None: continue
        kap = str(row[1]).strip() if row[1] else '-'
        pusk = str(row[2]).strip()
        sp_ages=[si(row[3]),si(row[4]),si(row[5]),si(row[6]),si(row[7]),si(row[8])]
        po_ages=[si(row[10]),si(row[11]),si(row[12]),si(row[13]),si(row[14]),si(row[15])]
        ts=si(row[9]); tp=si(row[16]); r=sf(row[17])
        puskesmas.append(pusk); kapanewon.append(kap)
        total_suspek.append(ts); total_positif.append(tp); rate.append(r)
        for i,a in enumerate(age_labels):
            age_sp[a].append(sp_ages[i]); age_po[a].append(po_ages[i])
        kap_agg_sp[kap]=kap_agg_sp.get(kap,0)+ts
        kap_agg_po[kap]=kap_agg_po.get(kap,0)+tp
    tr = rows[31]
    age_sp_total=[si(tr[3]),si(tr[4]),si(tr[5]),si(tr[6]),si(tr[7]),si(tr[8])]
    age_po_total=[si(tr[10]),si(tr[11]),si(tr[12]),si(tr[13]),si(tr[14]),si(tr[15])]
    return jsonify({
        'puskesmas':puskesmas,'kapanewon':kapanewon,
        'total_suspek':total_suspek,'total_positif':total_positif,'rate':rate,
        'age_labels':age_labels,'age_sp':age_sp,'age_po':age_po,
        'age_sp_total':age_sp_total,'age_po_total':age_po_total,
        'kapanewon_labels':list(kap_agg_sp.keys()),
        'kapanewon_suspek':list(kap_agg_sp.values()),
        'kapanewon_positif':list(kap_agg_po.values()),
        'grand_suspek':sum(total_suspek),'grand_positif':sum(total_positif)
    })

# ═══════════════════════════════════════════════════════════════
# CATATAN PIMPINAN — CRUD
# ═══════════════════════════════════════════════════════════════
@app.route('/api/catatan', methods=['GET'])
def catatan_list():
    prioritas = request.args.get('prioritas')
    with get_db() as conn:
        if prioritas and prioritas != 'semua':
            rows = conn.execute(
                'SELECT * FROM catatan WHERE prioritas=? ORDER BY dibuat DESC', (prioritas,)
            ).fetchall()
        else:
            rows = conn.execute(
                'SELECT * FROM catatan ORDER BY dibuat DESC'
            ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/catatan', methods=['POST'])
def catatan_create():
    d = request.json
    now = datetime.datetime.now().strftime('%d %b %Y, %H:%M')
    with get_db() as conn:
        cur = conn.execute(
            'INSERT INTO catatan (judul,isi,prioritas,bidang,penulis,dibuat) VALUES (?,?,?,?,?,?)',
            (d['judul'], d['isi'], d.get('prioritas','Sedang'),
             d.get('bidang','Umum'), d.get('penulis','Kepala Dinas'), now)
        )
        conn.commit()
        row = conn.execute('SELECT * FROM catatan WHERE id=?', (cur.lastrowid,)).fetchone()
    return jsonify(dict(row)), 201

@app.route('/api/catatan/<int:cid>', methods=['PUT'])
def catatan_update(cid):
    d = request.json
    now = datetime.datetime.now().strftime('%d %b %Y, %H:%M')
    with get_db() as conn:
        conn.execute(
            'UPDATE catatan SET judul=?,isi=?,prioritas=?,bidang=?,diubah=? WHERE id=?',
            (d['judul'], d['isi'], d.get('prioritas','Sedang'),
             d.get('bidang','Umum'), now, cid)
        )
        conn.commit()
        row = conn.execute('SELECT * FROM catatan WHERE id=?', (cid,)).fetchone()
    return jsonify(dict(row))

@app.route('/api/catatan/<int:cid>', methods=['DELETE'])
def catatan_delete(cid):
    with get_db() as conn:
        conn.execute('DELETE FROM catatan WHERE id=?', (cid,))
        conn.commit()
    return jsonify({'deleted': cid})

# ═══════════════════════════════════════════════════════════════
# ADMIN INPUT — DBD / STUNTING / CAMPAK
# ═══════════════════════════════════════════════════════════════
def init_admin_db():
    with get_db() as conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS admin_dbd (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                tahun       TEXT, kapanewon TEXT, puskesmas TEXT,
                usia_0      INTEGER DEFAULT 0, usia_1  INTEGER DEFAULT 0,
                usia_2      INTEGER DEFAULT 0, usia_3  INTEGER DEFAULT 0,
                usia_4      INTEGER DEFAULT 0, usia_5  INTEGER DEFAULT 0,
                usia_6      INTEGER DEFAULT 0, usia_7  INTEGER DEFAULT 0,
                laki        INTEGER DEFAULT 0, perempuan INTEGER DEFAULT 0,
                meninggal   INTEGER DEFAULT 0,
                dibuat      TEXT
            )''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS admin_stunting (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                tahun       TEXT, puskesmas TEXT,
                sasaran     INTEGER DEFAULT 0, dipantau  INTEGER DEFAULT 0,
                sangat_pendek INTEGER DEFAULT 0, pendek  INTEGER DEFAULT 0,
                normal      INTEGER DEFAULT 0, tinggi   INTEGER DEFAULT 0,
                stunting    INTEGER DEFAULT 0,
                dibuat      TEXT
            )''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS admin_campak (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                tahun       TEXT, kapanewon TEXT, puskesmas TEXT,
                suspek      INTEGER DEFAULT 0, positif  INTEGER DEFAULT 0,
                dibuat      TEXT
            )''')
        conn.commit()

init_admin_db()

@app.route('/api/admin/dbd', methods=['GET','POST'])
def admin_dbd():
    if request.method == 'POST':
        d = request.json
        now = datetime.datetime.now().strftime('%d %b %Y, %H:%M')
        ages = d.get('ages', [0]*8)
        with get_db() as conn:
            conn.execute(
                '''INSERT INTO admin_dbd
                   (tahun,kapanewon,puskesmas,usia_0,usia_1,usia_2,usia_3,usia_4,usia_5,usia_6,usia_7,laki,perempuan,meninggal,dibuat)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (d.get('tahun'), d.get('kapanewon'), d.get('puskesmas'),
                 ages[0],ages[1],ages[2],ages[3],ages[4],ages[5],ages[6],ages[7],
                 d.get('laki',0), d.get('perempuan',0), d.get('meninggal',0), now)
            )
            conn.commit()
        return jsonify({'status': 'ok'}), 201
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM admin_dbd ORDER BY id DESC').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/admin/dbd/<int:rid>', methods=['DELETE'])
def admin_dbd_delete(rid):
    with get_db() as conn:
        conn.execute('DELETE FROM admin_dbd WHERE id=?', (rid,))
        conn.commit()
    return jsonify({'deleted': rid})

@app.route('/api/admin/stunting', methods=['GET','POST'])
def admin_stunting():
    if request.method == 'POST':
        d = request.json
        now = datetime.datetime.now().strftime('%d %b %Y, %H:%M')
        with get_db() as conn:
            conn.execute(
                '''INSERT INTO admin_stunting
                   (tahun,puskesmas,sasaran,dipantau,sangat_pendek,pendek,normal,tinggi,stunting,dibuat)
                   VALUES (?,?,?,?,?,?,?,?,?,?)''',
                (d.get('tahun'), d.get('puskesmas'),
                 d.get('sasaran',0), d.get('dipantau',0), d.get('sangat_pendek',0),
                 d.get('pendek',0), d.get('normal',0), d.get('tinggi',0), d.get('stunting',0), now)
            )
            conn.commit()
        return jsonify({'status': 'ok'}), 201
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM admin_stunting ORDER BY id DESC').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/admin/stunting/<int:rid>', methods=['DELETE'])
def admin_stunting_delete(rid):
    with get_db() as conn:
        conn.execute('DELETE FROM admin_stunting WHERE id=?', (rid,))
        conn.commit()
    return jsonify({'deleted': rid})

@app.route('/api/admin/campak', methods=['GET','POST'])
def admin_campak():
    if request.method == 'POST':
        d = request.json
        now = datetime.datetime.now().strftime('%d %b %Y, %H:%M')
        with get_db() as conn:
            conn.execute(
                '''INSERT INTO admin_campak
                   (tahun,kapanewon,puskesmas,suspek,positif,dibuat)
                   VALUES (?,?,?,?,?,?)''',
                (d.get('tahun'), d.get('kapanewon'), d.get('puskesmas'),
                 d.get('suspek',0), d.get('positif',0), now)
            )
            conn.commit()
        return jsonify({'status': 'ok'}), 201
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM admin_campak ORDER BY id DESC').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/admin/campak/<int:rid>', methods=['DELETE'])
def admin_campak_delete(rid):
    with get_db() as conn:
        conn.execute('DELETE FROM admin_campak WHERE id=?', (rid,))
        conn.commit()
    return jsonify({'deleted': rid})

@app.route('/')
def index(): return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True, port=5000)