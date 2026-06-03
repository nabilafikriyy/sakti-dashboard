from flask import Flask, jsonify, render_template, request
from flask_cors import CORS
import openpyxl, os, datetime
import mysql.connector
from mysql.connector import Error

app = Flask(__name__)
CORS(app)
DATA_DIR = os.path.join(os.path.dirname(__file__), 'data')

# ── Konfigurasi MySQL ──────────────────────────────────────────────
DB_CONFIG = {
    'host':     'localhost',
    'user':     'root',        # ganti sesuai user MySQL kamu
    'password': '',            # ganti sesuai password MySQL kamu
    'database': 'sakti_db',   # pastikan database ini sudah dibuat
    'charset':  'utf8mb4'
}

FILE_DBD      = 'analisis_dbd_per_kelompok_umur_updated.xlsx'
FILE_STUNTING = 'gizi_balita_psg_2023_2025_updated.xlsx'
FILE_CAMPAK   = 'template_campak_sleman_2023-2025.xlsx'

# ── INIT DATABASE ──────────────────────────────────────────────────
def get_db():
    conn = mysql.connector.connect(**DB_CONFIG)
    return conn

def init_db():
    try:
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS catatan (
                id        INT          AUTO_INCREMENT PRIMARY KEY,
                judul     VARCHAR(255) NOT NULL,
                isi       TEXT         NOT NULL,
                prioritas VARCHAR(50)  DEFAULT 'Sedang',
                bidang    VARCHAR(50)  DEFAULT 'Umum',
                penulis   VARCHAR(100) DEFAULT 'Kepala Dinas',
                dibuat    VARCHAR(50)  NOT NULL,
                diubah    VARCHAR(50)
            )
        ''')
        conn.commit()
        cursor.close()
        conn.close()
    except Error as e:
        print(f"Error init_db: {e}")

init_db()

def load_wb(f): return openpyxl.load_workbook(os.path.join(DATA_DIR, f), data_only=True)
def si(v): return int(v or 0)
def sf(v): return round(float(v or 0), 2)
def avg(lst): return round(sum(lst)/len(lst), 2) if lst else 0

# ═══════════════════════════════════════════════════════════════
# DBD
# ═══════════════════════════════════════════════════════════════
# Kolom updated DBD year sheet:
# 0=No, 1=Kapanewon, 2=Puskesmas, 3-10=Kel.Umur(<1,1-4,5-9,10-14,15-19,20-44,45-59,≥60)
# 11=TOTAL(formula→None), 12=LAKI-LAKI, 13=PEREMPUAN, 14=MENINGGAL, 15=CFR

def _read_dbd_year(tahun):
    """Baca satu sheet tahun DBD, kembalikan dict per puskesmas."""
    ws = load_wb(FILE_DBD)[str(tahun)]
    rows = list(ws.iter_rows(values_only=True))
    result = {}
    for row in rows[5:30]:
        if row[2] is None: continue
        pusk = str(row[2]).strip()
        kap  = str(row[1]).strip() if row[1] else '-'
        ages = [si(row[i]) for i in range(3, 11)]
        total = sum(ages)
        result[pusk] = {
            'kapanewon': kap,
            'ages': ages,
            'total': total,
            'laki': si(row[12]),
            'perempuan': si(row[13]),
            'meninggal': si(row[14]),
            'cfr': sf(row[15]),
        }
    return result

@app.route('/api/dbd/rekapitulasi')
def dbd_rekapitulasi():
    data = {}
    for yr in [2023, 2024, 2025]:
        data[yr] = _read_dbd_year(yr)

    # Urutan puskesmas dari sheet 2023 (urutan asli)
    ws_order = load_wb(FILE_DBD)['2023']
    rows_order = list(ws_order.iter_rows(values_only=True))
    puskesmas = [str(r[2]).strip() for r in rows_order[5:30] if r[2]]

    y23 = [data[2023].get(p, {}).get('total', 0) for p in puskesmas]
    y24 = [data[2024].get(p, {}).get('total', 0) for p in puskesmas]
    y25 = [data[2025].get(p, {}).get('total', 0) for p in puskesmas]

    return jsonify({
        'puskesmas': puskesmas,
        'tahun_2023': y23, 'tahun_2024': y24, 'tahun_2025': y25,
        'total_2023': sum(y23), 'total_2024': sum(y24), 'total_2025': sum(y25),
        'laki_2023':  sum(data[2023].get(p,{}).get('laki',0) for p in puskesmas),
        'perempuan_2023': sum(data[2023].get(p,{}).get('perempuan',0) for p in puskesmas),
        'laki_2024':  sum(data[2024].get(p,{}).get('laki',0) for p in puskesmas),
        'perempuan_2024': sum(data[2024].get(p,{}).get('perempuan',0) for p in puskesmas),
        'laki_2025':  sum(data[2025].get(p,{}).get('laki',0) for p in puskesmas),
        'perempuan_2025': sum(data[2025].get(p,{}).get('perempuan',0) for p in puskesmas),
    })

@app.route('/api/dbd/kelompok_umur')
def dbd_kelompok_umur():
    age_labels = ['<1','1-4','5-9','10-14','15-19','20-44','45-59','≥60']
    result = {}
    for yr in [2023, 2024, 2025]:
        d = _read_dbd_year(yr)
        totals = [0]*8
        for pusk_data in d.values():
            for i, v in enumerate(pusk_data['ages']):
                totals[i] += v
        result[yr] = totals
    return jsonify({
        'age_groups': age_labels,
        'data_2023': result[2023],
        'data_2024': result[2024],
        'data_2025': result[2025],
    })

@app.route('/api/dbd/gender')
def dbd_gender():
    result = {}
    for yr in [2023, 2024, 2025]:
        d = _read_dbd_year(yr)
        result[yr] = {
            'laki': sum(v['laki'] for v in d.values()),
            'perempuan': sum(v['perempuan'] for v in d.values()),
        }
    return jsonify({
        'laki_2023': result[2023]['laki'], 'perempuan_2023': result[2023]['perempuan'],
        'laki_2024': result[2024]['laki'], 'perempuan_2024': result[2024]['perempuan'],
        'laki_2025': result[2025]['laki'], 'perempuan_2025': result[2025]['perempuan'],
    })

@app.route('/api/dbd/pertahun/<int:tahun>')
def dbd_pertahun(tahun):
    ws = load_wb(FILE_DBD)[str(tahun)]
    rows = list(ws.iter_rows(values_only=True))
    age_labels = ['<1','1-4','5-9','10-14','15-19','20-44','45-59','≥60']
    puskesmas, kapanewon, total, meninggal, cfr = [], [], [], [], []
    laki_list, perempuan_list = [], []
    age_data = {a: [] for a in age_labels}
    kapanewon_agg = {}

    for row in rows[5:30]:
        if row[2] is None: continue
        kap  = str(row[1]).strip() if row[1] else '-'
        pusk = str(row[2]).strip()
        ages = [si(row[i]) for i in range(3, 11)]
        tot  = sum(ages)
        men  = si(row[14])
        c    = sf(row[15])
        l    = si(row[12])
        p    = si(row[13])

        puskesmas.append(pusk); kapanewon.append(kap)
        total.append(tot); meninggal.append(men); cfr.append(c)
        laki_list.append(l); perempuan_list.append(p)
        for i, a in enumerate(age_labels): age_data[a].append(ages[i])
        kapanewon_agg[kap] = kapanewon_agg.get(kap, 0) + tot

    age_total = [sum(age_data[a]) for a in age_labels]

    return jsonify({
        'puskesmas': puskesmas, 'kapanewon': kapanewon,
        'total': total, 'meninggal': meninggal, 'cfr': cfr,
        'laki': laki_list, 'perempuan': perempuan_list,
        'age_data': age_data, 'age_labels': age_labels, 'age_total': age_total,
        'kapanewon_labels': list(kapanewon_agg.keys()),
        'kapanewon_total': list(kapanewon_agg.values()),
        'grand_total': sum(total), 'total_meninggal': sum(meninggal),
    })

# ═══════════════════════════════════════════════════════════════
# STUNTING
# ═══════════════════════════════════════════════════════════════
# Kolom updated Stunting year sheet (0-indexed):
# 0=No, 1=Puskesmas, 2=Sasaran, 3=Dipantau,
# 4=SP_n, 5=SP%, 6=P_n, 7=P%, 8=N_n, 9=N%, 10=T_n, 11=T%,
# 12=Stunted_n, 13=Stunted%, 14=Stunting_n, 15=Stunting%
# 16=Cakupan% (formula→None, hitung manual)

def _read_stunting_year(tahun):
    ws = load_wb(FILE_STUNTING)[str(tahun)]
    rows = list(ws.iter_rows(values_only=True))
    puskesmas, sasaran, dipantau = [], [], []
    sp_n, sp_pct, p_n, p_pct = [], [], [], []
    normal_n, normal_pct, tinggi_n, tinggi_pct = [], [], [], []
    stunted_n, stunted_pct, stunting_n, stunting_pct, cakupan = [], [], [], [], []

    for row in rows[6:31]:
        if row[1] is None or str(row[1]).strip().upper() in ('JUMLAH','TOTAL',''): continue
        sas = si(row[2]); dip = si(row[3])
        puskesmas.append(str(row[1]).strip())
        sasaran.append(sas); dipantau.append(dip)
        sp_n.append(si(row[4]));    sp_pct.append(sf(row[5]))
        p_n.append(si(row[6]));     p_pct.append(sf(row[7]))
        normal_n.append(si(row[8]));  normal_pct.append(sf(row[9]))
        tinggi_n.append(si(row[10])); tinggi_pct.append(sf(row[11]))
        stunted_n.append(si(row[12]));  stunted_pct.append(sf(row[13]))
        stunting_n.append(si(row[14])); stunting_pct.append(sf(row[15]))
        # Cakupan: hitung manual karena kolom formula = None
        cak = round(dip / sas * 100, 2) if sas > 0 else 0
        cakupan.append(cak)

    return {
        'puskesmas': puskesmas, 'sasaran': sasaran, 'dipantau': dipantau,
        'sangat_pendek_n': sp_n, 'sangat_pendek_pct': sp_pct,
        'pendek_n': p_n, 'pendek_pct': p_pct,
        'normal_n': normal_n, 'normal_pct': normal_pct,
        'tinggi_n': tinggi_n, 'tinggi_pct': tinggi_pct,
        'stunted_n': stunted_n, 'stunted_pct': stunted_pct,
        'stunting_n': stunting_n, 'stunting_pct': stunting_pct,
        'cakupan': cakupan,
        # Kabupaten aggregate (untuk donut)
        'kab_sp': sum(sp_n), 'kab_p': sum(p_n),
        'kab_normal': sum(normal_n), 'kab_tinggi': sum(tinggi_n),
        # Total row (hitung dari data)
        'total_sasaran': sum(sasaran), 'total_dipantau': sum(dipantau),
        'total_stunted': sum(stunted_n), 'total_stunting': sum(stunting_n),
        'avg_cakupan': avg(cakupan),
    }

@app.route('/api/stunting/rekapitulasi')
def stunting_rekapitulasi():
    d23 = _read_stunting_year(2023)
    d24 = _read_stunting_year(2024)
    d25 = _read_stunting_year(2025)

    # Gunakan puskesmas dari 2023 sebagai urutan acuan
    puskesmas = d23['puskesmas']

    def get_pct(d, key, p):
        if p in d['puskesmas']:
            i = d['puskesmas'].index(p)
            return d[key][i]
        return 0.0

    s23 = [get_pct(d23, 'stunting_pct', p) for p in puskesmas]
    s24 = [get_pct(d24, 'stunting_pct', p) for p in puskesmas]
    s25 = [get_pct(d25, 'stunting_pct', p) for p in puskesmas]

    return jsonify({
        'puskesmas': puskesmas,
        'stunting_2023': s23, 'stunting_2024': s24, 'stunting_2025': s25,
        'avg_2023': avg(s23), 'avg_2024': avg(s24), 'avg_2025': avg(s25),
    })

@app.route('/api/stunting/pertahun/<int:tahun>')
def stunting_pertahun(tahun):
    d = _read_stunting_year(tahun)
    return jsonify(d)

# ═══════════════════════════════════════════════════════════════
# CAMPAK (tidak berubah, file tidak dimodifikasi)
# ═══════════════════════════════════════════════════════════════
@app.route('/api/campak/rekapitulasi')
def campak_rekapitulasi():
    ws = load_wb(FILE_CAMPAK)['Rekapitulasi']
    rows = list(ws.iter_rows(values_only=True))
    puskesmas=[]
    sp23,po23,r23,sp24,po24,r24,sp25,po25,r25=[],[],[],[],[],[],[],[],[]
    for row in rows[5:30]:
        if row[1] is None: continue
        puskesmas.append(str(row[1]).strip())
        sp23.append(si(row[2])); po23.append(si(row[3])); r23.append(sf(row[4]))
        sp24.append(si(row[5])); po24.append(si(row[6])); r24.append(sf(row[7]))
        sp25.append(si(row[8])); po25.append(si(row[9])); r25.append(sf(row[10] or 0))
    return jsonify({'puskesmas':puskesmas,
        'suspek_2023':sp23,'positif_2023':po23,'rate_2023':r23,
        'suspek_2024':sp24,'positif_2024':po24,'rate_2024':r24,
        'suspek_2025':sp25,'positif_2025':po25,'rate_2025':r25,
        'total_suspek_2023':sum(sp23),'total_positif_2023':sum(po23),
        'total_suspek_2024':sum(sp24),'total_positif_2024':sum(po24),
        'total_suspek_2025':sum(sp25),'total_positif_2025':sum(po25)})

@app.route('/api/campak/pertahun/<int:tahun>')
def campak_pertahun(tahun):
    ws = load_wb(FILE_CAMPAK)[f'Kasus {tahun}']
    rows = list(ws.iter_rows(values_only=True))
    age_labels = ['<1','1-4','5-9','10-14','15-19','≥20']
    puskesmas,kapanewon,total_suspek,total_positif,rate=[],[],[],[],[]
    kap_agg_sp,kap_agg_po={},{}
    age_sp={a:[] for a in age_labels}
    age_po={a:[] for a in age_labels}
    for row in rows[6:31]:
        if row[2] is None: continue
        kap = str(row[1]).strip() if row[1] else '-'
        pusk = str(row[2]).strip()
        sp_ages=[si(row[3]),si(row[4]),si(row[5]),si(row[6]),si(row[7]),si(row[8])]
        po_ages=[si(row[10]),si(row[11]),si(row[12]),si(row[13]),si(row[14]),si(row[15])]
        ts=si(row[9]); tp=si(row[16]); r=sf(row[17])
        puskesmas.append(pusk); kapanewon.append(kap)
        total_suspek.append(ts); total_positif.append(tp); rate.append(r)
        for i,a in enumerate(age_labels):
            age_sp[a].append(sp_ages[i]); age_po[a].append(po_ages[i])
        kap_agg_sp[kap]=kap_agg_sp.get(kap,0)+ts
        kap_agg_po[kap]=kap_agg_po.get(kap,0)+tp
    tr = rows[31]
    age_sp_total=[si(tr[3]),si(tr[4]),si(tr[5]),si(tr[6]),si(tr[7]),si(tr[8])]
    age_po_total=[si(tr[10]),si(tr[11]),si(tr[12]),si(tr[13]),si(tr[14]),si(tr[15])]
    return jsonify({
        'puskesmas':puskesmas,'kapanewon':kapanewon,
        'total_suspek':total_suspek,'total_positif':total_positif,'rate':rate,
        'age_labels':age_labels,'age_sp':age_sp,'age_po':age_po,
        'age_sp_total':age_sp_total,'age_po_total':age_po_total,
        'kapanewon_labels':list(kap_agg_sp.keys()),
        'kapanewon_suspek':list(kap_agg_sp.values()),
        'kapanewon_positif':list(kap_agg_po.values()),
        'grand_suspek':sum(total_suspek),'grand_positif':sum(total_positif)
    })

# ═══════════════════════════════════════════════════════════════
# CATATAN PIMPINAN — CRUD
# ═══════════════════════════════════════════════════════════════
@app.route('/api/catatan', methods=['GET'])
def catatan_list():
    prioritas = request.args.get('prioritas')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    if prioritas and prioritas != 'semua':
        cursor.execute('SELECT * FROM catatan WHERE prioritas=%s ORDER BY dibuat DESC', (prioritas,))
    else:
        cursor.execute('SELECT * FROM catatan ORDER BY dibuat DESC')
    rows = cursor.fetchall()
    cursor.close()
    conn.close()
    return jsonify(rows)

@app.route('/api/catatan', methods=['POST'])
def catatan_create():
    d = request.json
    now = datetime.datetime.now().strftime('%d %b %Y, %H:%M')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    cursor.execute(
        'INSERT INTO catatan (judul,isi,prioritas,bidang,penulis,dibuat) VALUES (%s,%s,%s,%s,%s,%s)',
        (d['judul'], d['isi'], d.get('prioritas','Sedang'),
         d.get('bidang','Umum'), d.get('penulis','Kepala Dinas'), now)
    )
    conn.commit()
    new_id = cursor.lastrowid
    cursor.execute('SELECT * FROM catatan WHERE id=%s', (new_id,))
    row = cursor.fetchone()
    cursor.close()
    conn.close()
    return jsonify(row), 201

@app.route('/api/catatan/<int:cid>', methods=['PUT'])
def catatan_update(cid):
    d = request.json
    now = datetime.datetime.now().strftime('%d %b %Y, %H:%M')
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    cursor.execute(
        'UPDATE catatan SET judul=%s,isi=%s,prioritas=%s,bidang=%s,diubah=%s WHERE id=%s',
        (d['judul'], d['isi'], d.get('prioritas','Sedang'),
         d.get('bidang','Umum'), now, cid)
    )
    conn.commit()
    cursor.execute('SELECT * FROM catatan WHERE id=%s', (cid,))
    row = cursor.fetchone()
    cursor.close()
    conn.close()
    return jsonify(row)

@app.route('/api/catatan/<int:cid>', methods=['DELETE'])
def catatan_delete(cid):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM catatan WHERE id=%s', (cid,))
    conn.commit()
    cursor.close()
    conn.close()
    return jsonify({'deleted': cid})

@app.route('/')
def index(): return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True, port=5000)
